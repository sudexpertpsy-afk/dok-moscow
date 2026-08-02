"""Персональные функции законодательства в кабинете (W-35)."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO

from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import (
    ActFragment,
    ActVersion,
    LawBookmark,
    LawNote,
    LawView,
    LawWatch,
    LawWatchNotice,
    LegalAct,
    TariffCode,
    User,
    utcnow,
)
from app.services.billing import get_tariff_limits
from app.services.legal_public import published_for
from app.services.legal_search import article_anchor, html_to_text
from app.services.mail import send_email

PAID_TARIFFS = frozenset({TariffCode.specialist, TariffCode.organization})
VIEW_HISTORY_LIMIT = 20
SOURCE_PORTAL = "Официальный интернет-портал правовой информации"


@dataclass
class PaidAccess:
    allowed: bool
    is_guest: bool
    tariff_code: TariffCode


def paid_access(db: Session, org_id: int | None) -> PaidAccess:
    if org_id is None:
        return PaidAccess(allowed=False, is_guest=True, tariff_code=TariffCode.guest)
    limits = get_tariff_limits(db, org_id)
    paid = limits.tariff_code in PAID_TARIFFS and limits.is_current
    return PaidAccess(
        allowed=paid,
        is_guest=limits.tariff_code == TariffCode.guest or not paid,
        tariff_code=limits.tariff_code,
    )


def bookmark_key(act_id: int, fragment_id: int | None) -> str:
    return f"{act_id}:{fragment_id or 0}"


def toggle_bookmark(
    db: Session,
    *,
    user_id: int,
    act_id: int,
    fragment_id: int | None = None,
) -> bool:
    """Добавить/снять закладку. Возвращает True если теперь в закладках."""
    key = bookmark_key(act_id, fragment_id)
    row = db.scalar(
        select(LawBookmark).where(
            LawBookmark.user_id == user_id, LawBookmark.bookmark_key == key
        )
    )
    if row is not None:
        db.delete(row)
        db.flush()
        return False
    db.add(
        LawBookmark(
            user_id=user_id,
            act_id=act_id,
            fragment_id=fragment_id,
            bookmark_key=key,
        )
    )
    db.flush()
    return True


def list_bookmarks_rich(db: Session, user_id: int) -> list[dict]:
    rows = db.scalars(
        select(LawBookmark)
        .where(LawBookmark.user_id == user_id)
        .order_by(LawBookmark.created_at.desc())
        .limit(100)
    ).all()
    out: list[dict] = []
    for b in rows:
        act = db.get(LegalAct, b.act_id)
        if act is None:
            continue
        frag = db.get(ActFragment, b.fragment_id) if b.fragment_id else None
        href = f"/cabinet/zakon/{act.slug}"
        if frag:
            href += f"#{article_anchor(frag.article_ref)}"
        out.append(
            {
                "act": act,
                "fragment": frag,
                "href": href,
                "label": (
                    f"{frag.article_ref} — {act.title}" if frag else act.title
                ),
            }
        )
    return out


def bookmarked_keys(db: Session, user_id: int, act_id: int) -> set[str]:
    rows = db.scalars(
        select(LawBookmark.bookmark_key).where(
            LawBookmark.user_id == user_id, LawBookmark.act_id == act_id
        )
    ).all()
    return set(rows)


def upsert_note(
    db: Session, *, user_id: int, fragment_id: int, body: str
) -> LawNote | None:
    text = (body or "").strip()
    row = db.scalar(
        select(LawNote).where(
            LawNote.user_id == user_id, LawNote.fragment_id == fragment_id
        )
    )
    if not text:
        if row is not None:
            db.delete(row)
            db.flush()
        return None
    if row is None:
        row = LawNote(user_id=user_id, fragment_id=fragment_id, body=text)
        db.add(row)
    else:
        row.body = text
        row.updated_at = utcnow()
    db.flush()
    return row


def notes_for_act(db: Session, user_id: int, act_id: int) -> dict[int, str]:
    frag_ids = list(
        db.scalars(select(ActFragment.id).where(ActFragment.act_id == act_id)).all()
    )
    if not frag_ids:
        return {}
    rows = db.scalars(
        select(LawNote).where(
            LawNote.user_id == user_id, LawNote.fragment_id.in_(frag_ids)
        )
    ).all()
    return {r.fragment_id: r.body for r in rows}


def toggle_watch(db: Session, *, user_id: int, act_id: int) -> bool:
    row = db.scalar(
        select(LawWatch).where(LawWatch.user_id == user_id, LawWatch.act_id == act_id)
    )
    if row is not None:
        db.delete(row)
        db.flush()
        return False
    db.add(LawWatch(user_id=user_id, act_id=act_id))
    db.flush()
    return True


def is_watching(db: Session, user_id: int, act_id: int) -> bool:
    return (
        db.scalar(
            select(LawWatch.id).where(
                LawWatch.user_id == user_id, LawWatch.act_id == act_id
            )
        )
        is not None
    )


def list_watches(db: Session, user_id: int) -> list[LegalAct]:
    act_ids = list(
        db.scalars(
            select(LawWatch.act_id)
            .where(LawWatch.user_id == user_id)
            .order_by(LawWatch.created_at.desc())
        ).all()
    )
    if not act_ids:
        return []
    acts = {
        a.id: a
        for a in db.scalars(select(LegalAct).where(LegalAct.id.in_(act_ids))).all()
    }
    return [acts[i] for i in act_ids if i in acts]


def record_view(
    db: Session,
    *,
    user_id: int,
    act_id: int,
    fragment_id: int | None = None,
) -> None:
    db.add(
        LawView(
            user_id=user_id,
            act_id=act_id,
            fragment_id=fragment_id,
            viewed_at=utcnow(),
        )
    )
    db.flush()


def recent_views(db: Session, user_id: int, *, limit: int = VIEW_HISTORY_LIMIT) -> list[dict]:
    rows = db.scalars(
        select(LawView)
        .where(LawView.user_id == user_id)
        .order_by(LawView.viewed_at.desc())
        .limit(200)
    ).all()
    seen: set[int] = set()
    out: list[dict] = []
    for v in rows:
        if v.act_id in seen:
            continue
        seen.add(v.act_id)
        act = db.get(LegalAct, v.act_id)
        if act is None:
            continue
        out.append({"act": act, "viewed_at": v.viewed_at, "href": f"/cabinet/zakon/{act.slug}"})
        if len(out) >= limit:
            break
    return out


def quote_for_fragment(act: LegalAct, fragment: ActFragment | None, published: ActVersion | None) -> str:
    """Цитата с реквизитом для вставки в заключение."""
    rev = ""
    if published and published.revision_date:
        rev = f" (ред. от {published.revision_date.strftime('%d.%m.%Y')})"
    elif published:
        rev = " (опубликованная редакция)"
    num = (act.number or "").strip() or act.title
    if fragment and fragment.article_ref:
        head = f"{fragment.article_ref} {num}{rev}"
        body = html_to_text(fragment.body_html or "")
    else:
        head = f"{num}{rev}"
        body = html_to_text((published.body_html if published else "") or "")
    body = re.sub(r"\s+", " ", body).strip()
    if len(body) > 1200:
        body = body[:1197].rstrip() + "…"
    return f"{head} // {SOURCE_PORTAL}\n\n{body}".strip()


def export_fragments_docx(
    act: LegalAct,
    fragments: list[ActFragment],
    published: ActVersion | None,
) -> bytes:
    """DOCX-извлечение статей с шапкой редакции."""
    doc = DocxDocument()
    rev = "без даты"
    if published and published.revision_date:
        rev = published.revision_date.strftime("%d.%m.%Y")
    doc.add_heading("Извлечение из нормативного акта", level=1)
    doc.add_paragraph(f"Извлечение из {act.title}")
    if act.number:
        doc.add_paragraph(f"Реквизиты: {act.number}")
    doc.add_paragraph(f"Редакция от {rev}")
    doc.add_paragraph(f"Источник: {SOURCE_PORTAL}")
    doc.add_paragraph("")
    for frag in fragments:
        title = frag.article_ref
        if frag.title and frag.title != frag.article_ref:
            title = f"{frag.article_ref} — {frag.title}"
        doc.add_heading(title, level=2)
        text = html_to_text(frag.body_html or "") or "Текст статьи отсутствует."
        for para in re.split(r"\n+", text):
            p = para.strip()
            if p:
                doc.add_paragraph(p)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def queue_watch_notices(db: Session, act_id: int, version_id: int) -> int:
    """После публикации — поставить уведомления подписчикам в очередь."""
    user_ids = list(
        db.scalars(select(LawWatch.user_id).where(LawWatch.act_id == act_id)).all()
    )
    n = 0
    for uid in user_ids:
        db.add(
            LawWatchNotice(
                user_id=uid,
                act_id=act_id,
                version_id=version_id,
            )
        )
        n += 1
    db.flush()
    return n


def _day_start(now: datetime | None = None) -> datetime:
    now = now or utcnow()
    if now.tzinfo is None:
        now = now.replace(tzinfo=timezone.utc)
    return datetime(now.year, now.month, now.day, tzinfo=timezone.utc)


def user_digest_sent_today(db: Session, user_id: int, *, now: datetime | None = None) -> bool:
    start = _day_start(now)
    return (
        db.scalar(
            select(LawWatchNotice.id).where(
                LawWatchNotice.user_id == user_id,
                LawWatchNotice.sent_at.is_not(None),
                LawWatchNotice.sent_at >= start,
            ).limit(1)
        )
        is not None
    )


def flush_watch_digests(db: Session, *, force_user_id: int | None = None) -> int:
    """Отправить дайджесты: не более одного письма в сутки на пользователя.

    Если сегодня уже отправляли — оставляем notice в очереди до следующих суток.
    """
    settings = get_settings()
    q = select(LawWatchNotice.user_id).where(LawWatchNotice.sent_at.is_(None)).distinct()
    if force_user_id is not None:
        q = q.where(LawWatchNotice.user_id == force_user_id)
    user_ids = list(db.scalars(q).all())
    sent_count = 0
    app = settings.app_base_url.rstrip("/")
    for uid in user_ids:
        if user_digest_sent_today(db, uid):
            continue
        notices = list(
            db.scalars(
                select(LawWatchNotice)
                .where(LawWatchNotice.user_id == uid, LawWatchNotice.sent_at.is_(None))
                .order_by(LawWatchNotice.created_at.asc())
            ).all()
        )
        if not notices:
            continue
        user = db.get(User, uid)
        if user is None or not (user.email or "").strip():
            continue
        lines = [
            "Опубликованы новые редакции нормативных актов, на которые вы подписаны:",
            "",
        ]
        for n in notices:
            act = db.get(LegalAct, n.act_id)
            ver = db.get(ActVersion, n.version_id)
            if act is None:
                continue
            rev = ""
            if ver and ver.revision_date:
                rev = f" (ред. от {ver.revision_date.strftime('%d.%m.%Y')})"
            lines.append(f"• {act.title}{rev}")
            lines.append(f"  Текст: {app}/cabinet/zakon/{act.slug}")
            lines.append(f"  Diff в админке: {app}/admin/legal/{act.id}")
            lines.append("")
        lines.append("Отключить слежение можно на странице акта в кабинете.")
        ok = send_email(
            settings,
            to_addr=user.email,
            subject=f"[Док.Москва] Изменения НПА ({len(notices)})",
            body="\n".join(lines),
        )
        if ok or not settings.smtp_host:
            # без SMTP считаем «отправлено» (лог в send_email) — чтобы тесты и dev не копили очередь
            now = utcnow()
            for n in notices:
                n.sent_at = now
            sent_count += 1
            db.flush()
    return sent_count


def notify_watchers_after_publish(db: Session, act_id: int, version_id: int) -> int:
    queued = queue_watch_notices(db, act_id, version_id)
    if queued:
        flush_watch_digests(db)
    return queued


def cabinet_hit_url(slug: str, article_ref: str = "") -> str:
    base = f"/cabinet/zakon/{slug}"
    if article_ref:
        return f"{base}#{article_anchor(article_ref)}"
    return base
