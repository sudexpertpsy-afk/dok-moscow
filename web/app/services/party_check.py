"""Проверка контрагента через DaData findById (W-21)."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models import (
    Counterparty,
    CounterpartySource,
    CounterpartyType,
    PartyCheck,
    TariffCode,
    utcnow,
)
from app.services.billing import ensure_payment_settings, get_tariff_limits
from app.services.counterparties import apply_fields
from app.services.dadata import (
    PartyCard,
    find_party,
    party_card_to_counterparty_fields,
    party_check_usage_today,
    suggest,
)
from app.services.gotenberg import GotenbergError, convert_docx_to_pdf

PAID_TARIFFS = {TariffCode.specialist, TariffCode.organization}
STALE_DAYS = 30
INACTIVE_STATUSES = {"LIQUIDATING", "LIQUIDATED", "BANKRUPT"}


@dataclass
class FieldDiff:
    field: str
    label: str
    old: str
    new: str


@dataclass
class AccessState:
    allowed: bool
    is_guest: bool
    tariff_code: TariffCode
    limit: int
    used: int
    remaining: int


def party_check_daily_limit(db: Session) -> int:
    row = ensure_payment_settings(db)
    return int(row.party_check_daily_limit or 100)


def access_state(db: Session, org_id: int) -> AccessState:
    limits = get_tariff_limits(db, org_id)
    limit = party_check_daily_limit(db)
    used = party_check_usage_today(db, org_id)
    paid = limits.tariff_code in PAID_TARIFFS and limits.is_current
    remaining = max(0, limit - used)
    return AccessState(
        allowed=paid and remaining > 0,
        is_guest=limits.tariff_code == TariffCode.guest or not paid,
        tariff_code=limits.tariff_code,
        limit=limit,
        used=used,
        remaining=remaining,
    )


def can_use_party_check(db: Session, org_id: int) -> bool:
    state = access_state(db, org_id)
    return (not state.is_guest) and state.tariff_code in PAID_TARIFFS and state.remaining >= 0


def has_paid_access(db: Session, org_id: int) -> bool:
    limits = get_tariff_limits(db, org_id)
    return limits.tariff_code in PAID_TARIFFS and limits.is_current


def is_digits_query(query: str) -> bool:
    q = (query or "").strip()
    digits = "".join(ch for ch in q if ch.isdigit())
    return bool(q) and digits == q.replace(" ", "").replace("-", "") and len(digits) >= 10


def search_parties(
    db: Session,
    *,
    org_id: int,
    user_id: int | None,
    query: str,
) -> tuple[PartyCard | None, list[dict[str, str]], str | None]:
    """Цифры → findById; текст → suggest/party. Возвращает (card, candidates, error)."""
    q = (query or "").strip()
    if not q:
        return None, [], "Введите ИНН, ОГРН или название"

    state = access_state(db, org_id)
    if state.is_guest or state.tariff_code not in PAID_TARIFFS:
        return None, [], "Раздел доступен на тарифах «Специалист» и «Организация»"
    if state.used >= state.limit:
        return None, [], f"Исчерпан суточный лимит проверок ({state.limit})"

    if is_digits_query(q):
        card = find_party(db, org_id=org_id, user_id=user_id, query="".join(ch for ch in q if ch.isdigit()))
        if card is None:
            return None, [], "Организация не найдена в ЕГРЮЛ/ЕГРИП"
        return card, [], None

    items = suggest(db, org_id=org_id, user_id=user_id, kind="party", query=q, count=10)
    candidates = []
    for item in items:
        data = item.data or {}
        name = data.get("name") or {}
        title = str(
            name.get("short_with_opf")
            or name.get("full_with_opf")
            or item.value
            or ""
        )
        inn = str(data.get("inn") or "")
        if not inn:
            continue
        status = str((data.get("state") or {}).get("status") or "")
        candidates.append(
            {
                "inn": inn,
                "name": title,
                "ogrn": str(data.get("ogrn") or ""),
                "address": str(
                    (data.get("address") or {}).get("value")
                    if isinstance(data.get("address"), dict)
                    else ""
                ),
                "status": status,
            }
        )
    if not candidates:
        return None, [], "Ничего не найдено. Уточните название или укажите ИНН."
    return None, candidates, None


def load_party_by_inn(
    db: Session,
    *,
    org_id: int,
    user_id: int | None,
    inn: str,
    query: str | None = None,
    record_journal: bool = True,
) -> tuple[PartyCard | None, PartyCheck | None, str | None]:
    state = access_state(db, org_id)
    if state.is_guest or state.tariff_code not in PAID_TARIFFS:
        return None, None, "Раздел доступен на тарифах «Специалист» и «Организация»"
    if state.used >= state.limit:
        return None, None, f"Исчерпан суточный лимит проверок ({state.limit})"

    digits = "".join(ch for ch in (inn or "") if ch.isdigit())
    if len(digits) not in (10, 12):
        return None, None, "Некорректный ИНН"

    card = find_party(db, org_id=org_id, user_id=user_id, query=digits)
    if card is None:
        return None, None, "Организация не найдена в ЕГРЮЛ/ЕГРИП"

    check = None
    if record_journal:
        check = record_party_check(
            db,
            org_id=org_id,
            user_id=user_id,
            card=card,
            query=query or digits,
        )
    return card, check, None


def record_party_check(
    db: Session,
    *,
    org_id: int,
    user_id: int | None,
    card: PartyCard,
    query: str,
    counterparty_id: int | None = None,
) -> PartyCheck:
    if counterparty_id is None and card.inn:
        cp = db.scalar(
            select(Counterparty).where(
                Counterparty.org_id == org_id,
                Counterparty.inn == card.inn,
            )
        )
        if cp is not None:
            counterparty_id = cp.id
            cp.egrul_status = card.status
            cp.egrul_checked_at = utcnow()

    row = PartyCheck(
        org_id=org_id,
        user_id=user_id,
        counterparty_id=counterparty_id,
        inn=card.inn or "",
        query=(query or "")[:255],
        status=card.status,
        status_label=card.status_label,
        checked_at=utcnow(),
        snapshot=card.raw or {},
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def find_counterparty_by_inn(db: Session, org_id: int, inn: str) -> Counterparty | None:
    digits = "".join(ch for ch in (inn or "") if ch.isdigit())
    if not digits:
        return None
    return db.scalar(
        select(Counterparty).where(Counterparty.org_id == org_id, Counterparty.inn == digits)
    )


DIFF_LABELS = {
    "name": "Наименование",
    "fio": "Руководитель / ФИО",
    "inn": "ИНН",
    "kpp": "КПП",
    "ogrn": "ОГРН",
    "address": "Адрес",
}


def diff_counterparty(cp: Counterparty, card: PartyCard) -> list[FieldDiff]:
    fields = party_card_to_counterparty_fields(card)
    out: list[FieldDiff] = []
    for key, label in DIFF_LABELS.items():
        old = str(getattr(cp, key) or "").strip()
        new = str(fields.get(key) or "").strip()
        if old != new:
            out.append(FieldDiff(field=key, label=label, old=old, new=new))
    return out


def upsert_counterparty_from_card(
    db: Session,
    *,
    org_id: int,
    card: PartyCard,
) -> tuple[Counterparty, list[FieldDiff], bool]:
    """Создать или обновить контрагента. Возвращает (cp, diff, created)."""
    fields = party_card_to_counterparty_fields(card)
    existing = find_counterparty_by_inn(db, org_id, card.inn)
    if existing is None:
        cp = Counterparty(
            org_id=org_id,
            type=CounterpartyType.ul if card.party_type == "LEGAL" else CounterpartyType.fl,
            source=CounterpartySource.dadata,
        )
        apply_fields(cp, fields, source=CounterpartySource.dadata)
        cp.egrul_status = card.status
        cp.egrul_checked_at = utcnow()
        db.add(cp)
        db.commit()
        db.refresh(cp)
        return cp, [], True

    diffs = diff_counterparty(existing, card)
    apply_fields(existing, fields, source=CounterpartySource.dadata)
    existing.egrul_status = card.status
    existing.egrul_checked_at = utcnow()
    db.commit()
    db.refresh(existing)
    return existing, diffs, False


def list_party_checks(
    db: Session,
    org_id: int,
    *,
    counterparty_id: int | None = None,
    inn: str | None = None,
    limit: int = 100,
) -> list[PartyCheck]:
    stmt = (
        select(PartyCheck)
        .options(joinedload(PartyCheck.user), joinedload(PartyCheck.counterparty))
        .where(PartyCheck.org_id == org_id)
        .order_by(PartyCheck.checked_at.desc())
        .limit(limit)
    )
    if counterparty_id is not None:
        stmt = stmt.where(PartyCheck.counterparty_id == counterparty_id)
    if inn:
        digits = "".join(ch for ch in inn if ch.isdigit())
        if digits:
            stmt = stmt.where(PartyCheck.inn == digits)
    return list(db.scalars(stmt).unique().all())


def get_party_check(db: Session, org_id: int, check_id: int) -> PartyCheck | None:
    row = db.get(PartyCheck, check_id)
    if row is None or row.org_id != org_id:
        return None
    return row


def refresh_stale_egrul(
    db: Session,
    *,
    org_id: int,
    user_id: int | None,
    cp: Counterparty,
) -> tuple[bool, str | None]:
    """Обновить статус ЕГРЮЛ, если старше 30 дней. Возвращает (updated, warning)."""
    if cp.type != CounterpartyType.ul or not cp.inn:
        return False, None

    checked = cp.egrul_checked_at
    if checked is not None:
        if checked.tzinfo is None:
            checked = checked.replace(tzinfo=timezone.utc)
        if utcnow() - checked < timedelta(days=STALE_DAYS):
            warning = None
            if (cp.egrul_status or "").upper() in INACTIVE_STATUSES or (
                cp.egrul_status and cp.egrul_status.upper() != "ACTIVE"
            ):
                if (cp.egrul_status or "").upper() != "ACTIVE":
                    warning = (
                        f"Статус контрагента в ЕГРЮЛ: «{cp.egrul_status or 'не действует'}». "
                        "Проверьте карточку до генерации договора."
                    )
            return False, warning

    state = access_state(db, org_id)
    if state.is_guest or state.used >= state.limit:
        if cp.egrul_status and cp.egrul_status.upper() != "ACTIVE":
            return False, (
                f"Статус контрагента в ЕГРЮЛ: «{cp.egrul_status}». "
                "Проверьте карточку до генерации договора."
            )
        return False, None

    card = find_party(db, org_id=org_id, user_id=user_id, query=cp.inn)
    if card is None:
        return False, None

    cp.egrul_status = card.status
    cp.egrul_checked_at = utcnow()
    record_party_check(
        db,
        org_id=org_id,
        user_id=user_id,
        card=card,
        query=cp.inn,
        counterparty_id=cp.id,
    )
    warning = None
    if card.status and card.status.upper() != "ACTIVE":
        warning = (
            f"Статус контрагента в ЕГРЮЛ: «{card.status_label or card.status}». "
            "Проверьте карточку до генерации договора."
        )
    return True, warning


def build_party_card_pdf(card: PartyCard, checked_at: datetime | None = None) -> bytes:
    """Печатная карточка: DOCX → PDF (Gotenberg), иначе текстовый PDF."""
    checked_at = checked_at or utcnow()
    docx_buf = _party_card_docx(card, checked_at)
    try:
        from pathlib import Path
        import tempfile

        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "party_check.docx"
            pdf_path = Path(tmp) / "party_check.pdf"
            docx_path.write_bytes(docx_buf)
            convert_docx_to_pdf(docx_path, pdf_path)
            return pdf_path.read_bytes()
    except (GotenbergError, OSError, Exception):
        return _simple_party_pdf(card, checked_at)


def _party_card_docx(card: PartyCard, checked_at: datetime) -> bytes:
    doc = DocxDocument()
    doc.add_heading("Карточка проверки контрагента", level=1)
    doc.add_paragraph(f"Дата проверки: {checked_at.strftime('%d.%m.%Y %H:%M')} UTC")
    doc.add_paragraph("Источник: ЕГРЮЛ через сервис DaData")
    doc.add_paragraph(f"Статус на дату: {card.status_label or card.status or '—'}")
    doc.add_paragraph("")
    for label, value in card.display_rows():
        doc.add_paragraph(f"{label}: {value}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_escape(text: str) -> str:
    return (
        text.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )


def _to_pdf_latin(text: str) -> str:
    """ASCII-safe строка для Helvetica (кириллица → транслит/пропуск)."""
    table = str.maketrans(
        {
            "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e",
            "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
            "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
            "ф": "f", "х": "h", "ц": "c", "ч": "ch", "ш": "sh", "щ": "sch",
            "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
            "А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Е": "E", "Ё": "E",
            "Ж": "Zh", "З": "Z", "И": "I", "Й": "Y", "К": "K", "Л": "L", "М": "M",
            "Н": "N", "О": "O", "П": "P", "Р": "R", "С": "S", "Т": "T", "У": "U",
            "Ф": "F", "Х": "H", "Ц": "C", "Ч": "Ch", "Ш": "Sh", "Щ": "Sch",
            "Ъ": "", "Ы": "Y", "Ь": "", "Э": "E", "Ю": "Yu", "Я": "Ya",
            "«": '"', "»": '"', "—": "-", "–": "-", "№": "N",
        }
    )
    return "".join(ch if ord(ch) < 128 else "?" for ch in text.translate(table))


def _simple_party_pdf(card: PartyCard, checked_at: datetime) -> bytes:
    lines = [
        "Proverka kontragentа",
        f"Date: {checked_at.strftime('%d.%m.%Y %H:%M')} UTC",
        "Source: EGRUL via DaData",
        f"Status: {_to_pdf_latin(card.status_label or card.status or '-')}",
        f"INN: {card.inn or '-'}",
        f"OGRN: {card.ogrn or '-'}",
        f"Name: {_to_pdf_latin(card.name_short or card.name_full or '-')}",
    ]
    for label, value in card.display_rows()[:20]:
        lines.append(f"{_to_pdf_latin(label)}: {_to_pdf_latin(value)}")

    y = 800
    content_cmds = ["BT", "/F1 11 Tf", "50 800 Td", "14 TL"]
    for i, line in enumerate(lines):
        esc = _pdf_escape(_to_pdf_latin(line)[:110])
        if i == 0:
            content_cmds.append(f"({esc}) Tj")
        else:
            content_cmds.append(f"T* ({esc}) Tj")
    content_cmds.append("ET")
    stream = "\n".join(content_cmds).encode("latin-1", errors="replace")

    objects: list[bytes] = [
        b"1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n",
        b"2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n",
        (
            b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>endobj\n"
        ),
        (
            f"4 0 obj<< /Length {len(stream)} >>stream\n".encode("ascii")
            + stream
            + b"\nendstream\nendobj\n"
        ),
        b"5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj\n",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(out))
        out.extend(obj)
    xref_pos = len(out)
    out.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    out.extend(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        out.extend(f"{off:010d} 00000 n \n".encode("ascii"))
    out.extend(
        (
            f"trailer<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(out)
