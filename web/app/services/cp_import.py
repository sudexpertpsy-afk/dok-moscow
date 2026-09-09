"""W-49 A: импорт контрагентов из CSV/XLSX и journal.jsonl."""

from __future__ import annotations

import csv
import io
import json
import re
import uuid
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from enum import StrEnum
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import (
    Counterparty,
    CounterpartySource,
    CounterpartyType,
    TariffCode,
    utcnow,
)
from app.services.audit import record_event
from app.services.billing import get_tariff_limits
from app.services.counterparties import apply_fields, validate_counterparty_form
from app.services.rate_counters import get_count, incr_counter, set_if_absent
from app.services.safe_paths import org_files_root, resolve_under

MAX_FILE_BYTES = 5 * 1024 * 1024
SYNC_MAX_ROWS = 500
RATE_LIMIT_PER_HOUR = 10
GUEST_ROW_LIMIT = 50
PAID_ROW_LIMIT = 5000

FIELD_KEYS: tuple[str, ...] = (
    "type",
    "name",
    "fio",
    "inn",
    "kpp",
    "ogrn",
    "snils",
    "passport_series",
    "passport_number",
    "passport_issuer",
    "passport_date",
    "address",
    "phone",
    "email",
    "bank_name",
    "bank_bik",
    "bank_account",
    "bank_corr_account",
    "notes",
)

FIELD_LABELS: dict[str, str] = {
    "type": "Тип",
    "name": "Наименование",
    "fio": "ФИО",
    "inn": "ИНН",
    "kpp": "КПП",
    "ogrn": "ОГРН",
    "snils": "СНИЛС",
    "passport_series": "Паспорт серия",
    "passport_number": "Паспорт номер",
    "passport_issuer": "Кем выдан",
    "passport_date": "Дата выдачи",
    "address": "Адрес",
    "phone": "Телефон",
    "email": "E-mail",
    "bank_name": "Банк",
    "bank_bik": "БИК",
    "bank_account": "Расчётный счёт",
    "bank_corr_account": "Корр. счёт",
    "notes": "Примечание",
}

_SYNONYMS: dict[str, tuple[str, ...]] = {
    "type": ("тип", "type", "вид"),
    "name": (
        "наименование",
        "название",
        "организация",
        "название организации",
        "название_заказчика",
        "name",
    ),
    "fio": (
        "фио",
        "ф.и.о.",
        "ф.и.о",
        "фамилия имя отчество",
        "фио_клиента",
        "фио_эксперта",
    ),
    "inn": ("инн", "inn", "инн_заказчика"),
    "kpp": ("кпп", "kpp", "кпп_заказчика"),
    "ogrn": ("огрн", "огрнип", "ogrn", "огрн_заказчика"),
    "snils": ("снилс", "snils"),
    "passport_series": ("паспорт серия", "серия паспорта", "паспорт_серия", "серия"),
    "passport_number": ("паспорт номер", "номер паспорта", "паспорт_номер", "номер"),
    "passport_issuer": ("кем выдан", "паспорт_выдан", "выдан"),
    "passport_date": ("дата выдачи", "паспорт_дата"),
    "address": ("адрес", "юр_адрес", "адрес_клиента", "адрес регистрации"),
    "phone": ("телефон", "тел.", "тел", "phone", "телефон_клиента"),
    "email": ("email", "e-mail", "почта", "эл. почта", "email_клиента"),
    "bank_name": ("банк", "название банка"),
    "bank_bik": ("бик", "bik"),
    "bank_account": ("расчётный счёт", "расчетный счет", "счёт", "счет", "р/с"),
    "bank_corr_account": ("корр. счёт", "корр счет", "к/с", "коррсчёт"),
    "notes": ("примечание", "заметки", "комментарий"),
}


class ImportErrorMsg(ValueError):
    """Понятная ошибка импорта для UI."""


class DuplicateMode(StrEnum):
    skip = "skip"
    fill = "fill"
    overwrite = "overwrite"


class RowStatus(StrEnum):
    ok = "ok"
    warn = "warn"
    error = "error"


@dataclass
class ParsedTable:
    headers: list[str]
    rows: list[list[str]]
    source_kind: str  # csv / xlsx / journal


@dataclass
class PreviewRow:
    index: int
    status: RowStatus
    messages: list[str]
    typ: CounterpartyType
    fields: dict[str, str | None]
    duplicate: str | None = None
    intra_merge: bool = False


@dataclass
class PreviewResult:
    rows: list[PreviewRow]
    ready: int
    warnings: int
    errors: int
    enrich_needed: int
    dadata_used: int
    dadata_limit: int


@dataclass
class CommitResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: int = 0
    report_rel: str | None = None


def row_limit_for_org(db: Session, org_id: int) -> int:
    limits = get_tariff_limits(db, org_id)
    if limits.tariff_code == TariffCode.guest:
        return GUEST_ROW_LIMIT
    return PAID_ROW_LIMIT


def _hour_window_start() -> datetime:
    now = utcnow()
    return now.replace(minute=0, second=0, microsecond=0)


def _rate_key(org_id: int) -> str:
    return f"cp_import:{org_id}:{_hour_window_start().strftime('%Y%m%d%H')}"


def check_import_rate(db: Session, org_id: int) -> None:
    key = _rate_key(org_id)
    set_if_absent(db, key, window_start=_hour_window_start(), count=0)
    if get_count(db, key) >= RATE_LIMIT_PER_HOUR:
        raise ImportErrorMsg("Слишком много импортов. Подождите час (лимит 10 в час).")


def consume_import_rate(db: Session, org_id: int) -> None:
    key = _rate_key(org_id)
    set_if_absent(db, key, window_start=_hour_window_start(), count=0)
    n = incr_counter(db, key, window_start=_hour_window_start(), by=1)
    if n > RATE_LIMIT_PER_HOUR:
        raise ImportErrorMsg("Слишком много импортов. Подождите час (лимит 10 в час).")


def imports_dir(org_id: int) -> Path:
    d = org_files_root(org_id) / "imports"
    d.mkdir(parents=True, exist_ok=True)
    return d


def token_dir(org_id: int, token: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        raise ImportErrorMsg("Некорректный сеанс импорта.")
    return resolve_under(imports_dir(org_id), token)


def new_token() -> str:
    return uuid.uuid4().hex


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def normalize_header(raw: str) -> str:
    s = (raw or "").strip().lower().replace("ё", "е")
    s = re.sub(r"\s+", " ", s)
    return s


def auto_map_columns(headers: list[str]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = {k: None for k in FIELD_KEYS}
    used: set[int] = set()
    norms = [normalize_header(h) for h in headers]
    for key, syns in _SYNONYMS.items():
        for i, n in enumerate(norms):
            if i in used:
                continue
            if n in syns or n.replace(".", "") in syns:
                mapping[key] = i
                used.add(i)
                break
    return mapping


def detect_csv_dialect(sample: str) -> csv.Dialect:
    try:
        return csv.Sniffer().sniff(sample, delimiters=";\t,")
    except csv.Error:
        class _D(csv.Dialect):
            delimiter = ";" if sample.count(";") >= sample.count(",") else ","
            quotechar = '"'
            doublequote = True
            skipinitialspace = True
            lineterminator = "\n"
            quoting = csv.QUOTE_MINIMAL

        return _D()


def decode_bytes(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = raw.decode(enc)
        except UnicodeDecodeError:
            continue
        if enc == "cp1251" and "\ufffd" in text:
            continue
        # utf-8-sig уже снимает BOM; на всякий случай
        if text.startswith("\ufeff"):
            text = text.lstrip("\ufeff")
        return text
    raise ImportErrorMsg(
        "Не удалось прочитать файл. Сохраните CSV в кодировке UTF-8 или Windows-1251."
    )


def parse_csv_bytes(raw: bytes) -> ParsedTable:
    text = decode_bytes(raw)
    sample = text[:4096]
    dialect = detect_csv_dialect(sample)
    reader = csv.reader(io.StringIO(text), dialect)
    all_rows = [[cell_str(c) for c in row] for row in reader]
    all_rows = [r for r in all_rows if any(x.strip() for x in r)]
    if not all_rows:
        raise ImportErrorMsg("Файл пустой.")
    headers = all_rows[0]
    data = all_rows[1:]
    return ParsedTable(headers=headers, rows=data, source_kind="csv")


def parse_xlsx_bytes(raw: bytes) -> ParsedTable:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows: list[list[str]] = []
        for excel_row in ws.iter_rows(values_only=True):
            cells = [cell_str(c) for c in excel_row]
            if any(x.strip() for x in cells):
                rows.append(cells)
    finally:
        wb.close()
    if not rows:
        raise ImportErrorMsg("Таблица пустая.")
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return ParsedTable(headers=rows[0], rows=rows[1:], source_kind="xlsx")


def _journal_fields_from_context(ctx: dict) -> tuple[CounterpartyType, dict[str, str | None]] | None:
    name = cell_str(ctx.get("название_заказчика") or ctx.get("название"))
    inn = cell_str(ctx.get("инн_заказчика") or ctx.get("инн"))
    expert = cell_str(ctx.get("фио_эксперта"))
    fio = cell_str(ctx.get("фио_клиента") or ctx.get("фио"))
    if name or (inn and len(re.sub(r"\D", "", inn)) == 10):
        typ = CounterpartyType.ul
        fields = {
            "name": name or None,
            "inn": inn or None,
            "kpp": cell_str(ctx.get("кпп_заказчика") or ctx.get("кпп")) or None,
            "ogrn": cell_str(ctx.get("огрн_заказчика") or ctx.get("огрн")) or None,
            "address": cell_str(ctx.get("адрес_заказчика") or ctx.get("юр_адрес")) or None,
            "phone": cell_str(ctx.get("телефон_заказчика") or ctx.get("телефон")) or None,
            "email": cell_str(ctx.get("email_заказчика") or ctx.get("email")) or None,
        }
        return typ, fields
    if expert:
        return CounterpartyType.expert, {
            "fio": expert,
            "inn": inn or None,
            "address": cell_str(ctx.get("адрес_эксперта") or ctx.get("адрес")) or None,
        }
    if fio:
        return CounterpartyType.fl, {
            "fio": fio,
            "inn": inn or None,
            "passport_series": cell_str(ctx.get("паспорт_серия")) or None,
            "passport_number": cell_str(ctx.get("паспорт_номер")) or None,
            "passport_issuer": cell_str(ctx.get("паспорт_выдан")) or None,
            "address": cell_str(ctx.get("адрес_клиента") or ctx.get("адрес")) or None,
            "phone": cell_str(ctx.get("телефон_клиента") or ctx.get("телефон")) or None,
            "email": cell_str(ctx.get("email_клиента") or ctx.get("email")) or None,
        }
    return None


def parse_journal_bytes(raw: bytes) -> ParsedTable:
    text = decode_bytes(raw)
    extracted: list[dict[str, str]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        ctx = rec.get("context") if isinstance(rec, dict) else None
        if not isinstance(ctx, dict):
            continue
        parsed = _journal_fields_from_context(ctx)
        if parsed is None:
            continue
        typ, fields = parsed
        row = {k: "" for k in FIELD_KEYS}
        row["type"] = typ.value
        for k, v in fields.items():
            if v:
                row[k] = v
        extracted.append(row)
    if not extracted:
        raise ImportErrorMsg("В журнале не нашлось контрагентов (проверьте .journal.jsonl Шаблонера).")
    headers = list(FIELD_KEYS)
    rows = [[r.get(k, "") for k in FIELD_KEYS] for r in extracted]
    return ParsedTable(headers=headers, rows=rows, source_kind="journal")


def parse_upload(filename: str, raw: bytes) -> ParsedTable:
    name = (filename or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise ImportErrorMsg("Формат .xls не поддерживается. Сохраните файл как .xlsx или CSV.")
    if name.endswith(".jsonl"):
        return parse_journal_bytes(raw)
    if name.endswith(".xlsx"):
        return parse_xlsx_bytes(raw)
    if name.endswith((".csv", ".txt")):
        return parse_csv_bytes(raw)
    # эвристика
    if raw[:2] == b"PK":
        return parse_xlsx_bytes(raw)
    if raw.lstrip()[:1] == b"{":
        return parse_journal_bytes(raw)
    return parse_csv_bytes(raw)


def parse_type_cell(raw: str, inn: str, default: CounterpartyType | None) -> CounterpartyType:
    s = normalize_header(raw).replace(" ", "")
    if s in ("ul", "юл", "юрлицо", "организация", "ооо", "ип"):
        return CounterpartyType.ul
    if s in ("fl", "фл", "физлицо", "физическоелицо"):
        return CounterpartyType.fl
    if s in ("expert", "эксперт"):
        return CounterpartyType.expert
    digits = re.sub(r"\D", "", inn or "")
    if len(digits) == 10:
        return CounterpartyType.ul
    if len(digits) == 12:
        return CounterpartyType.fl
    return default or CounterpartyType.fl


def mapped_fields(row: list[str], mapping: dict[str, int | None]) -> dict[str, str]:
    out: dict[str, str] = {}
    for key in FIELD_KEYS:
        idx = mapping.get(key)
        if idx is None or idx < 0 or idx >= len(row):
            out[key] = ""
        else:
            out[key] = cell_str(row[idx])
    return out


def _digits(s: str | None) -> str:
    return re.sub(r"\D", "", s or "")


def _norm_fio(s: str | None) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def duplicate_key(typ: CounterpartyType, fields: dict[str, str | None]) -> tuple | None:
    if typ == CounterpartyType.ul:
        inn = _digits(fields.get("inn"))
        if not inn:
            return None
        return ("ul", inn, _digits(fields.get("kpp")))
    if typ == CounterpartyType.expert:
        fio = _norm_fio(fields.get("fio"))
        inn = _digits(fields.get("inn"))
        if fio and inn:
            return ("expert", fio, inn)
        return None
    fio = _norm_fio(fields.get("fio"))
    ser = _digits(fields.get("passport_series"))
    num = _digits(fields.get("passport_number"))
    if fio and (ser or num):
        return ("fl", fio, ser, num)
    return None


def _cp_key(cp: Counterparty) -> tuple | None:
    fields = {k: getattr(cp, k, None) for k in FIELD_KEYS if k != "type"}
    return duplicate_key(cp.type, fields)


def _to_store(raw: dict[str, str]) -> dict[str, str | None]:
    return {k: (v.strip() or None) for k, v in raw.items() if k != "type"}


def merge_empty(base: dict[str, str | None], extra: dict[str, str | None]) -> dict[str, str | None]:
    out = dict(base)
    for k, v in extra.items():
        if v and not out.get(k):
            out[k] = v
    return out


def classify_rows(
    table: ParsedTable,
    mapping: dict[str, int | None],
    *,
    default_type: CounterpartyType | None,
    existing: list[Counterparty],
) -> list[PreviewRow]:
    existing_map: dict[tuple, Counterparty] = {}
    for cp in existing:
        key = _cp_key(cp)
        if key:
            existing_map[key] = cp

    seen: dict[tuple, int] = {}
    out: list[PreviewRow] = []
    for i, row in enumerate(table.rows):
        raw = mapped_fields(row, mapping)
        typ = parse_type_cell(raw.get("type") or "", raw.get("inn") or "", default_type)
        fields = _to_store(raw)
        msgs: list[str] = []
        errs = validate_counterparty_form(typ, fields)
        msgs.extend(errs)
        key = duplicate_key(typ, fields)
        intra = False
        dup_label = None
        if key is not None and key in seen:
            intra = True
            prev = out[seen[key]]
            prev.fields = merge_empty(prev.fields, fields)
            prev.intra_merge = True
            prev.messages.append("Объединено с дублем в файле")
            if prev.status == RowStatus.ok:
                prev.status = RowStatus.warn
            continue
        if key is not None:
            seen[key] = len(out)
            if key in existing_map:
                dup_label = "есть в картотеке"
        if typ == CounterpartyType.ul and not fields.get("inn"):
            msgs.append("Нет ИНН — карточка сохранится без проверки ЕГРЮЛ")
        status = RowStatus.ok
        if errs:
            status = RowStatus.error
        elif msgs or dup_label or intra:
            status = RowStatus.warn
        out.append(
            PreviewRow(
                index=i + 2,
                status=status,
                messages=msgs,
                typ=typ,
                fields=fields,
                duplicate=dup_label,
                intra_merge=intra,
            )
        )
    return out


def summarize(rows: list[PreviewRow], *, dadata_used: int, dadata_limit: int) -> PreviewResult:
    ready = sum(1 for r in rows if r.status == RowStatus.ok)
    warnings = sum(1 for r in rows if r.status == RowStatus.warn)
    errors = sum(1 for r in rows if r.status == RowStatus.error)
    enrich = sum(
        1
        for r in rows
        if r.typ == CounterpartyType.ul
        and _digits(r.fields.get("inn"))
        and not (r.fields.get("name") or "").strip()
        and r.status != RowStatus.error
    )
    return PreviewResult(
        rows=rows,
        ready=ready,
        warnings=warnings,
        errors=errors,
        enrich_needed=enrich,
        dadata_used=dadata_used,
        dadata_limit=dadata_limit,
    )


def _apply_mode(cp: Counterparty, fields: dict[str, str | None], mode: DuplicateMode) -> None:
    if mode == DuplicateMode.overwrite:
        apply_fields(cp, fields)
        return
    if mode == DuplicateMode.fill:
        patch = {k: v for k, v in fields.items() if v and not getattr(cp, k, None)}
        if patch:
            apply_fields(cp, patch)


def commit_rows(
    db: Session,
    org_id: int,
    rows: Iterable[PreviewRow],
    *,
    mode: DuplicateMode,
    require_zero_errors: bool,
    user_id: int | None,
    filename: str,
    token: str | None = None,
    enrich: bool = False,
) -> CommitResult:
    rows = list(rows)
    if require_zero_errors and any(r.status == RowStatus.error for r in rows):
        raise ImportErrorMsg("В файле есть ошибки. Исправьте их или снимите «только при 0 ошибок».")

    existing = list(db.scalars(select(Counterparty).where(Counterparty.org_id == org_id)).all())
    by_key: dict[tuple, Counterparty] = {}
    for cp in existing:
        key = _cp_key(cp)
        if key:
            by_key[key] = cp

    result = CommitResult()
    error_lines: list[tuple[int, str, dict]] = []
    batch = 0
    for prow in rows:
        if prow.status == RowStatus.error:
            result.errors += 1
            error_lines.append((prow.index, "; ".join(prow.messages), prow.fields))
            continue
        fields = dict(prow.fields)
        if enrich and prow.typ == CounterpartyType.ul and _digits(fields.get("inn")) and not fields.get("name"):
            try:
                from app.services.dadata import (
                    find_party,
                    party_card_to_counterparty_fields,
                )

                card = find_party(db, org_id=org_id, user_id=user_id, query=_digits(fields.get("inn")))
                if card is not None:
                    extra = party_card_to_counterparty_fields(card)
                    fields = merge_empty(fields, {k: extra.get(k) for k in fields})
                    fields["name"] = fields.get("name") or extra.get("name")
            except Exception as enrich_exc:  # noqa: BLE001
                # DaData недоступна — импорт без обогащения
                _ = enrich_exc
        key = duplicate_key(prow.typ, fields)
        found = by_key.get(key) if key else None
        if found is not None:
            if mode == DuplicateMode.skip:
                result.skipped += 1
                continue
            _apply_mode(found, fields, mode)
            result.updated += 1
        else:
            cp = Counterparty(
                org_id=org_id,
                type=prow.typ,
                source=CounterpartySource.manual,
            )
            apply_fields(cp, fields)
            db.add(cp)
            db.flush()
            result.created += 1
            new_key = _cp_key(cp)
            if new_key:
                by_key[new_key] = cp
        batch += 1
        if batch % 200 == 0:
            db.flush()

    report_rel = None
    if error_lines:
        report_rel = _write_error_report(org_id, error_lines, token=token)
        result.report_rel = report_rel

    record_event(
        db,
        type="counterparties.imported",
        org_id=org_id,
        user_id=user_id,
        details={
            "filename": filename[:200],
            "created": result.created,
            "updated": result.updated,
            "skipped": result.skipped,
            "errors": result.errors,
        },
        commit=False,
    )
    db.flush()
    return result


def _write_error_report(
    org_id: int, lines: list[tuple[int, str, dict]], *, token: str | None
) -> str:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Ошибки"
    headers = ["строка", "причина"] + [FIELD_LABELS[k] for k in FIELD_KEYS if k != "type"]
    ws.append(headers)
    for idx, reason, fields in lines:
        row = [idx, reason]
        for k in FIELD_KEYS:
            if k == "type":
                continue
            val = fields.get(k) or ""
            if str(val).startswith(("=", "+", "-", "@")):
                val = "'" + str(val)
            row.append(val)
        ws.append(row)
    dest = token_dir(org_id, token) if token else imports_dir(org_id) / new_token()
    dest.mkdir(parents=True, exist_ok=True)
    path = dest / "errors.xlsx"
    wb.save(path)
    root = Path(get_settings().files_root).resolve()
    return str(path.resolve().relative_to(root))


def build_template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    wb = Workbook()
    ws = wb.active
    ws.title = "Контрагенты"
    headers = [FIELD_LABELS[k] for k in FIELD_KEYS]
    ws.append(headers)
    example = {
        "type": "ЮЛ",
        "name": "ООО Ромашка",
        "inn": "7707083893",
        "kpp": "770701001",
        "ogrn": "1027700132195",
        "address": "г. Москва",
        "phone": "+7 495 000-00-00",
        "email": "office@example.com",
    }
    ws.append([example.get(k, "") for k in FIELD_KEYS])
    notes = {
        "type": "ФЛ / ЮЛ / Эксперт",
        "inn": "10 цифр ЮЛ или 12 ФЛ/ИП, с контрольной суммой",
        "kpp": "9 цифр",
        "ogrn": "13 или 15 цифр",
        "passport_date": "ДД.ММ.ГГГГ",
        "phone": "+7… или 8…",
    }
    for i, key in enumerate(FIELD_KEYS, start=1):
        hint = notes.get(key)
        if hint:
            ws.cell(1, i).comment = Comment(hint, "Док.Москва")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save_table_snapshot(org_id: int, token: str, table: ParsedTable, mapping: dict[str, int | None] | None = None) -> None:
    dest = token_dir(org_id, token)
    dest.mkdir(parents=True, exist_ok=True)
    payload = {
        "headers": table.headers,
        "rows": table.rows,
        "source_kind": table.source_kind,
        "mapping": mapping,
    }
    (dest / "table.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def load_table_snapshot(org_id: int, token: str) -> tuple[ParsedTable, dict[str, int | None] | None]:
    path = token_dir(org_id, token) / "table.json"
    if not path.is_file():
        raise ImportErrorMsg("Сеанс импорта не найден или истёк. Загрузите файл снова.")
    data = json.loads(path.read_text(encoding="utf-8"))
    table = ParsedTable(
        headers=list(data["headers"]),
        rows=list(data["rows"]),
        source_kind=str(data.get("source_kind") or "csv"),
    )
    mapping = data.get("mapping")
    return table, mapping


def save_mapping(
    org_id: int,
    token: str,
    mapping: dict[str, int | None],
    extra: dict[str, Any] | None = None,
) -> None:
    table, _ = load_table_snapshot(org_id, token)
    save_table_snapshot(org_id, token, table, mapping)
    meta = token_dir(org_id, token) / "meta.json"
    payload = extra or {}
    meta.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def load_meta(org_id: int, token: str) -> dict[str, Any]:
    path = token_dir(org_id, token) / "meta.json"
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def cleanup_stale_imports(*, older_than_hours: int = 24) -> int:
    import shutil

    root = Path(get_settings().files_root)
    if not root.is_dir():
        return 0
    cutoff = datetime.now(timezone.utc) - timedelta(hours=older_than_hours)
    removed = 0
    for org_dir in root.iterdir():
        imports = org_dir / "imports"
        if not imports.is_dir():
            continue
        for token_path in imports.iterdir():
            if not token_path.is_dir():
                continue
            mtime = datetime.fromtimestamp(token_path.stat().st_mtime, tz=timezone.utc)
            if mtime < cutoff:
                shutil.rmtree(token_path, ignore_errors=True)
                removed += 1
    return removed
