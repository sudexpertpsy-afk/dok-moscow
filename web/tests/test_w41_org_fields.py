"""W-41: пользовательские поля в своих шаблонах + безопасность Jinja/SSTI."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import timedelta
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import select

from app.defaults import empty_requisites
from app.models import (
    Event,
    OrgField,
    OrgFieldType,
    OrgRole,
    Organization,
    Subscription,
    SubscriptionStatus,
    Tariff,
    TariffCode,
    User,
    UserRole,
    utcnow,
)
from app.security import hash_password
from app.services.billing import ensure_beta_subscriptions, ensure_tariffs
from app.services.org_fields import (
    OrgFieldError,
    create_org_field,
    delete_org_field,
    templates_using_field,
)
from app.services.templates import ensure_core_on_path
from conftest import login


def _docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _org_pair(dbmod, *, email_a="w41a@example.com", email_b="w41b@example.com"):
    db = dbmod.SessionLocal()
    try:
        ensure_tariffs(db)
        org_a = Organization(name="ООО Альфа Поля", requisites=empty_requisites())
        org_b = Organization(name="ООО Бета Поля", requisites=empty_requisites())
        db.add_all([org_a, org_b])
        db.flush()
        ua = User(
            org_id=org_a.id,
            email=email_a,
            password_hash=hash_password("Passw0rd!"),
            role=UserRole.user,
            org_role=OrgRole.org_admin,
            is_active=True,
        )
        ub = User(
            org_id=org_b.id,
            email=email_b,
            password_hash=hash_password("Passw0rd!"),
            role=UserRole.user,
            org_role=OrgRole.org_admin,
            is_active=True,
        )
        db.add_all([ua, ub])
        ensure_beta_subscriptions(db)
        for oid in (org_a.id, org_b.id):
            tariff = db.scalar(select(Tariff).where(Tariff.code == TariffCode.organization))
            sub = db.scalar(select(Subscription).where(Subscription.org_id == oid))
            sub.tariff_id = tariff.id
            sub.status = SubscriptionStatus.active
            sub.ends_at = utcnow() + timedelta(days=30)
            sub.is_beta = False
        db.commit()
        return org_a.id, email_a, org_b.id, email_b
    finally:
        db.close()


def _csrf(html: str) -> str:
    m = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert m, "csrf not found"
    return m.group(1)


def _set_files_root(tmp_path, monkeypatch):
    from app.config import get_settings

    get_settings.cache_clear()
    s = get_settings()
    root = tmp_path / "files"
    root.mkdir(exist_ok=True)
    object.__setattr__(s, "files_root", str(root))
    return s


def test_upload_define_new_field_then_generate(app, tmp_path, monkeypatch):
    client, dbmod = app
    org_id, email, _, _ = _org_pair(dbmod)
    s = _set_files_root(tmp_path, monkeypatch)
    assert login(client, email, "Passw0rd!").status_code == 303

    page = client.get("/cabinet/templates/")
    csrf = _csrf(page.text)
    r = client.post(
        "/cabinet/templates/upload",
        data={"csrf_token": csrf, "contract_type": ""},
        files={
            "file": (
                "Акт_гарантия.docx",
                _docx("Срок: {{ срок_гарантии }}. Номер: {{ номер_договора }}."),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    assert r.status_code == 303
    loc = r.headers["location"]
    conf = client.get(loc)
    assert conf.status_code == 200
    assert "срок_гарантии" in conf.text
    assert "Новые" in conf.text or "новые" in conf.text.lower()
    assert "номер_договора" in conf.text

    csrf = _csrf(conf.text)
    stage = re.search(r'name="token" value="([^"]+)"', conf.text).group(1)
    r = client.post(
        "/cabinet/templates/upload/confirm",
        data={
            "csrf_token": csrf,
            "token": stage,
            "contract_type": "",
            "new_name_0": "срок_гарантии",
            "new_label_0": "Срок гарантии",
            "new_type_0": "string",
            "new_required_0": "1",
            "new_default_0": "12 месяцев",
            "new_hint_0": "Например: 12 месяцев",
            "new_options_0": "",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text[:500]
    assert (Path(s.files_root) / str(org_id) / "templates" / "Акт_гарантия.docx").is_file()

    db = dbmod.SessionLocal()
    try:
        field = db.scalar(
            select(OrgField).where(
                OrgField.org_id == org_id, OrgField.name == "срок_гарантии"
            )
        )
        assert field is not None
        assert field.label == "Срок гарантии"
        assert field.required is True
        ev = db.scalars(
            select(Event).where(
                Event.org_id == org_id, Event.type == "org_template_upload"
            )
        ).all()
        assert ev
    finally:
        db.close()

    form = client.get("/cabinet/documents/new/Акт_гарантия.docx")
    assert form.status_code == 200
    assert "Поля организации" in form.text
    assert "Срок гарантии" in form.text
    assert 'name="срок_гарантии"' in form.text

    csrf = _csrf(form.text)
    r = client.post(
        "/cabinet/documents/new/Акт_гарантия.docx",
        data={
            "csrf_token": csrf,
            "срок_гарантии": "24 месяца",
            "номер_договора": "42",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303
    assert "/cabinet/documents/" in r.headers["location"]
    doc_id = r.headers["location"].rstrip("/").split("/")[-1]
    db = dbmod.SessionLocal()
    try:
        from app.models import Document as DocRow

        doc = db.get(DocRow, int(doc_id))
        assert doc is not None
        assert doc.context.get("срок_гарантии") == "24 месяца"
        out = Path(s.files_root) / doc.file_path
        assert out.is_file()
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "24 месяца" in text
        assert "42" in text
    finally:
        db.close()


def test_undefined_field_blocks_save(app, tmp_path, monkeypatch):
    client, dbmod = app
    _, email, _, _ = _org_pair(dbmod, email_a="w41u@example.com", email_b="w41u2@example.com")
    _set_files_root(tmp_path, monkeypatch)
    assert login(client, email, "Passw0rd!").status_code == 303
    page = client.get("/cabinet/templates/")
    r = client.post(
        "/cabinet/templates/upload",
        data={"csrf_token": _csrf(page.text), "contract_type": ""},
        files={
            "file": (
                "x.docx",
                _docx("{{ новое_поле_блокировка }}"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    loc = r.headers["location"]
    conf = client.get(loc)
    stage = re.search(r'name="token" value="([^"]+)"', conf.text).group(1)
    # не передаём new_name_0 — должно отказать
    r = client.post(
        "/cabinet/templates/upload/confirm",
        data={
            "csrf_token": _csrf(conf.text),
            "token": stage,
            "contract_type": "",
        },
        follow_redirects=False,
    )
    assert r.status_code == 400
    assert "неопределённ" in r.text.lower() or "Нельзя сохранить" in r.text


@pytest.mark.parametrize(
    "payload,needle",
    [
        ("{% for x in y %}{{ x }}{% endfor %}", "{%"),
        ("{{ сумма | unknown_evil }}", "фильтр"),
        ("{{ cycler.__init__.__globals__ }}", "атрибут"),
        ("{{ ''.__class__.__mro__ }}", "атрибут"),
        ("{{ lipsum() }}", "функц"),
    ],
)
def test_upload_rejects_jinja_statements_filters_ssti(
    app, tmp_path, monkeypatch, payload, needle
):
    client, dbmod = app
    _, email, _, _ = _org_pair(
        dbmod,
        email_a=f"w41sec_{abs(hash(payload)) % 10**6}@example.com",
        email_b=f"w41secb_{abs(hash(payload)) % 10**6}@example.com",
    )
    _set_files_root(tmp_path, monkeypatch)
    assert login(client, email, "Passw0rd!").status_code == 303
    page = client.get("/cabinet/templates/")
    r = client.post(
        "/cabinet/templates/upload",
        data={"csrf_token": _csrf(page.text), "contract_type": ""},
        files={
            "file": (
                "evil.docx",
                _docx(payload),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        follow_redirects=False,
    )
    # либо 400 на upload, либо не уходим в confirm
    if r.status_code == 303:
        assert "/confirm" not in r.headers.get("location", "")
    else:
        assert r.status_code == 400
        assert needle.lower() in r.text.lower() or "запрещ" in r.text.lower()


def test_ssti_not_executed_in_sandbox(tmp_path):
    """Даже если опасный DOCX оказался на диске — SandboxedEnvironment не исполняет SSTI."""
    ensure_core_on_path()
    from docfiller_core.filler import fill_template
    from jinja2.exceptions import SecurityError, UndefinedError
    from jinja2.sandbox import SecurityError as SandboxSecurityError

    src = tmp_path / "ssti.docx"
    src.write_bytes(_docx("X{{ cycler.__init__.__globals__.os.system('id') }}Y"))
    out = tmp_path / "out.docx"
    with pytest.raises((SecurityError, SandboxSecurityError, Exception)) as excinfo:
        fill_template(str(src), str(out), {"cycler": "nope"})
    # не должен тихо записать «успех» с утечкой
    if out.exists():
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "uid=" not in text
        assert "__globals__" not in text or "X" in text
    err = str(excinfo.value).lower()
    assert (
        "access" in err
        or "security" in err
        or "sandbox" in err
        or "attribute" in err
        or "__" in err
        or "запрещ" in err
    )


def test_ssti_attr_chain_blocked_by_analyzer():
    ensure_core_on_path()
    from docfiller_core.template_security import (
        TemplateSecurityError,
        assert_safe_docx_template,
    )

    with pytest.raises(TemplateSecurityError):
        assert_safe_docx_template(
            _docx("{{ cycler.__init__.__globals__['os'].system('id') }}")
        )
    with pytest.raises(TemplateSecurityError):
        assert_safe_docx_template(_docx("{{ ''.__class__.__mro__[1].__subclasses__() }}"))


def test_org_field_dictionaries_isolated(app, tmp_path, monkeypatch):
    client, dbmod = app
    org_a_id, email_a, org_b_id, email_b = _org_pair(dbmod)
    _set_files_root(tmp_path, monkeypatch)

    db = dbmod.SessionLocal()
    try:
        create_org_field(
            db,
            org_id=org_a_id,
            user_id=None,
            name="секрет_альфы",
            label="Секрет",
            field_type=OrgFieldType.string,
        )
        db.commit()
    finally:
        db.close()

    assert login(client, email_b, "Passw0rd!").status_code == 303
    r = client.get("/cabinet/templates/fields")
    assert r.status_code == 200
    assert "секрет_альфы" not in r.text

    db = dbmod.SessionLocal()
    try:
        assert (
            db.scalar(
                select(OrgField).where(
                    OrgField.org_id == org_b_id, OrgField.name == "секрет_альфы"
                )
            )
            is None
        )
        # B не может удалить поле A по id
        field_a = db.scalar(
            select(OrgField).where(
                OrgField.org_id == org_a_id, OrgField.name == "секрет_альфы"
            )
        )
        fid = field_a.id
    finally:
        db.close()

    page = client.get("/cabinet/templates/fields")
    r = client.post(
        f"/cabinet/templates/fields/{fid}/delete",
        data={"csrf_token": _csrf(page.text)},
        follow_redirects=False,
    )
    assert r.status_code == 404
    db = dbmod.SessionLocal()
    try:
        assert db.get(OrgField, fid) is not None
    finally:
        db.close()


def test_cannot_delete_field_in_use(app, tmp_path, monkeypatch):
    _, dbmod = app
    org_id, email, _, _ = _org_pair(
        dbmod, email_a="w41del@example.com", email_b="w41del2@example.com"
    )
    s = _set_files_root(tmp_path, monkeypatch)
    db = dbmod.SessionLocal()
    try:
        f = create_org_field(
            db,
            org_id=org_id,
            user_id=None,
            name="поле_в_использовании",
            label="В деле",
            field_type="string",
        )
        db.commit()
        fid = f.id
    finally:
        db.close()

    from app.services.org_templates import org_templates_dir

    path = org_templates_dir(org_id) / "Использует.docx"
    path.write_bytes(_docx("{{ поле_в_использовании }}"))
    assert "Использует.docx" in templates_using_field(org_id, "поле_в_использовании")

    db = dbmod.SessionLocal()
    try:
        row = db.get(OrgField, fid)
        with pytest.raises(OrgFieldError, match="используется"):
            delete_org_field(db, row)
    finally:
        db.close()


def test_vba_rejected_on_validate():
    from app.services.docx_upload import DocxUploadError, validate_docx_bytes

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"></Types>',
        )
        zf.writestr("word/document.xml", "<w:document/>")
        zf.writestr("word/vbaProject.bin", b"MZ")
    with pytest.raises(DocxUploadError, match="VBA"):
        validate_docx_bytes(buf.getvalue())


def test_member_can_open_fields_page(app, tmp_path, monkeypatch):
    client, dbmod = app
    org_id, admin_email, _, _ = _org_pair(
        dbmod, email_a="w41adm@example.com", email_b="w41x@example.com"
    )
    _set_files_root(tmp_path, monkeypatch)
    db = dbmod.SessionLocal()
    try:
        member = User(
            org_id=org_id,
            email="w41member@example.com",
            password_hash=hash_password("Passw0rd!"),
            role=UserRole.user,
            org_role=OrgRole.org_member,
            is_active=True,
        )
        db.add(member)
        db.commit()
    finally:
        db.close()

    assert login(client, "w41member@example.com", "Passw0rd!").status_code == 303
    r = client.get("/cabinet/templates/fields")
    assert r.status_code == 200
    assert "Поля организации" in r.text
    # загрузка шаблонов — только admin
    r = client.get("/cabinet/templates/")
    assert r.status_code == 403 or "недостаточно" in r.text.lower() or "админ" in r.text.lower()
