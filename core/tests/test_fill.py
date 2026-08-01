"""Тесты рендера всех шаблонов (T-10 / test_fill)."""

from pathlib import Path

import pytest
from docx import Document

from docfiller_core.filler import fill_template, list_template_variables
from fixtures_sample_context import (
    TEMPLATES_DIR,
    build_context_for_template,
    expected_phrases,
    list_templates,
)


def _doc_full_text(path):
    """Текст документа: абзацы + ячейки таблиц."""
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return '\n'.join(parts)


@pytest.fixture
def templates_dir():
    assert TEMPLATES_DIR.is_dir(), f'Папка шаблонов не найдена: {TEMPLATES_DIR}'
    return TEMPLATES_DIR


@pytest.mark.parametrize('template_name', list_templates())
def test_list_template_variables(templates_dir, template_name):
    """Шаблон читается, переменные извлекаются без ошибок."""
    path = templates_dir / template_name
    variables = list_template_variables(path)
    assert isinstance(variables, list)
    # Системная переменная «настройки» не должна попадать в форму
    assert 'настройки' not in variables


@pytest.mark.parametrize('template_name', list_templates())
def test_render_template(templates_dir, template_name, tmp_path):
    """Каждый из 25 шаблонов рендерится без Jinja-ошибок."""
    src = templates_dir / template_name
    context = build_context_for_template(template_name)
    out = tmp_path / template_name

    fill_template(src, out, context)

    assert out.exists()
    assert out.stat().st_size > 500

    text = _doc_full_text(out)
    assert '{{' not in text, f'В документе остались неразобранные метки: {template_name}'
    assert '}}' not in text
    assert len(text.strip()) > 30, f'Слишком мало текста в {template_name}'

    for phrase in expected_phrases(template_name, context):
        assert phrase in text, (
            f'В «{template_name}» не найдена контрольная фраза «{phrase}»'
        )
