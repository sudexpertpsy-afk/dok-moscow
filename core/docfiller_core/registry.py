"""
Чтение/запись Excel-реестра документов.

Структура файла `Реестр.xlsx`:

    Лист "Реестр" — таблица документов
        Столбец A: Шаблон          (имя файла шаблона из папки «Шаблоны»)
        Столбец B: Файл_результат  (имя выходного файла; если пусто — будет сгенерировано)
        Столбец C: Статус          (заполняется после генерации: дата или ✗ ошибка)
        Столбцы D+: переменные шаблона (имя столбца = имя переменной)
"""

from pathlib import Path
from datetime import datetime
import shutil
import os
import openpyxl


HEADER_ROW = 1
DATA_START_ROW = 2

COL_TEMPLATE = 'Шаблон'
COL_OUTPUT   = 'Файл_результат'
COL_STATUS   = 'Статус'

SYSTEM_COLUMNS = {COL_TEMPLATE, COL_OUTPUT, COL_STATUS}


def is_locked_by_excel(registry_path):
    """Открыт ли файл реестра прямо сейчас в Excel?

    Excel создаёт рядом скрытый файл «~$Имя.xlsx». Также пробуем открыть
    файл на запись — это быстрая дополнительная проверка.
    """
    p = Path(registry_path)
    # Excel'овский lock-файл
    lock = p.parent / f'~${p.name}'
    if lock.exists():
        return True
    # Файл может быть открыт другим приложением — пробуем дописать в него
    try:
        with open(p, 'r+b'):
            pass
    except (PermissionError, OSError):
        return True
    return False


def backup_registry(registry_path):
    """Сохранить копию реестра в каталог бэкапов данных пользователя.

    Хранятся 10 последних резервных копий, более старые удаляются.
    """
    from . import paths

    p = Path(registry_path)
    if not p.exists():
        return None
    backup_dir = paths.backups_dir()
    backup_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    dst = backup_dir / f'{p.stem}_{ts}.xlsx'
    shutil.copy2(p, dst)

    # Чистка: оставляем 10 последних копий для этого реестра
    copies = sorted(
        backup_dir.glob(f'{p.stem}_*.xlsx'),
        key=lambda f: f.stat().st_mtime,
        reverse=True,
    )
    for old in copies[10:]:
        try:
            old.unlink()
        except OSError:
            pass
    return dst


def read_rows(registry_path):
    """
    Прочитать все строки реестра. Возвращает список словарей с ключами:
        '_row'      — номер строки в Excel,
        '_template' — имя шаблона,
        '_output'   — имя выходного файла (может быть пустым),
        '_status'   — текущий статус,
        остальные ключи — пользовательские поля (= имена столбцов).
    """
    # read_only=True — не блокирует файл, можно читать даже когда Excel открыт
    wb = openpyxl.load_workbook(registry_path, data_only=True, read_only=True)
    if 'Реестр' not in wb.sheetnames:
        wb.close()
        raise ValueError("В файле реестра не найден лист «Реестр»")
    ws = wb['Реестр']

    headers = {}
    rows_raw = list(ws.iter_rows(values_only=False))
    wb.close()
    if not rows_raw:
        return []

    # Заголовки из первой строки
    for col_idx, cell in enumerate(rows_raw[0], start=1):
        if cell.value:
            headers[col_idx] = str(cell.value).strip()

    rows = []
    for row_idx, row_cells in enumerate(rows_raw[1:], start=DATA_START_ROW):
        row_dict = {'_row': row_idx}
        has_data = False
        for col_idx, cell in enumerate(row_cells, start=1):
            header = headers.get(col_idx)
            if not header:
                continue
            value = cell.value
            if value is not None and value != '':
                has_data = True
            if header == COL_TEMPLATE:
                row_dict['_template'] = value
            elif header == COL_OUTPUT:
                row_dict['_output'] = value
            elif header == COL_STATUS:
                row_dict['_status'] = value
            else:
                row_dict[header] = value
        if has_data:
            rows.append(row_dict)
    return rows


def list_template_columns(registry_path):
    """Вернуть список пользовательских (не служебных) столбцов реестра."""
    wb = openpyxl.load_workbook(registry_path, data_only=True, read_only=True)
    if 'Реестр' not in wb.sheetnames:
        wb.close()
        return []
    ws = wb['Реестр']
    headers = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value and str(cell.value).strip() not in SYSTEM_COLUMNS:
            headers.append(str(cell.value).strip())
    wb.close()
    return headers


def write_status(registry_path, row_idx, output_filename, status):
    """Записать статус и имя выходного файла в строку реестра.

    Если Excel держит файл открытым — поднимаем PermissionError, который
    обрабатывает вызывающий код.
    """
    wb = openpyxl.load_workbook(registry_path)
    ws = wb['Реестр']

    # Найти индексы нужных столбцов по заголовкам
    target_cols = {}
    for col_idx, cell in enumerate(ws[HEADER_ROW], start=1):
        if cell.value in (COL_OUTPUT, COL_STATUS):
            target_cols[cell.value] = col_idx

    if COL_OUTPUT in target_cols and output_filename:
        ws.cell(row=row_idx, column=target_cols[COL_OUTPUT]).value = output_filename
    if COL_STATUS in target_cols:
        ws.cell(row=row_idx, column=target_cols[COL_STATUS]).value = status

    wb.save(registry_path)


def make_sample_registry(registry_path, sample_data, template_name='СППЭ_информация_суду.docx'):
    """Создать пример Реестр.xlsx с заголовками и одной строкой данных."""
    # TODO(этап 4+): vulture — не вызывается из UI; образец уже в Реестр_*.example.xlsx.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Реестр'

    headers = [COL_TEMPLATE, COL_OUTPUT, COL_STATUS] + list(sample_data.keys())
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = openpyxl.styles.Font(bold=True)

    row = [template_name, '', ''] + list(sample_data.values())
    for col_idx, val in enumerate(row, start=1):
        ws.cell(row=2, column=col_idx, value=val)

    # Подгон ширины колонок
    for col_idx, h in enumerate(headers, start=1):
        letter = openpyxl.utils.get_column_letter(col_idx)
        max_len = max(len(str(h)), len(str(row[col_idx - 1])) if col_idx - 1 < len(row) else 0)
        ws.column_dimensions[letter].width = min(max_len + 2, 60)

    # Вторая колонка («Файл_результат») чуть пошире
    ws.column_dimensions['B'].width = 28

    # Лист с инструкцией
    info = wb.create_sheet('Инструкция')
    info_lines = [
        'КАК ЗАПОЛНЯТЬ РЕЕСТР',
        '',
        '1. Каждая строка = один документ.',
        '2. В столбце «Шаблон» укажите имя файла из папки «Шаблоны».',
        '3. В столбце «Файл_результат» можно указать желаемое имя файла или оставить пустым.',
        '4. Имена остальных столбцов должны совпадать с переменными {{ ... }} в шаблоне.',
        '5. После генерации в столбце «Статус» появится дата создания.',
        '',
        'УМНЫЕ ФУНКЦИИ В ШАБЛОНАХ:',
        '   {{ "Лосев Сергей Александрович" | род }} → Лосева Сергея Александровича',
        '   {{ "Громинская Александра Владимировна" | дат }} → Громинской Александре Владимировне',
        '   {{ 250000 | руб }} → 250 000 (двести пятьдесят тысяч) рублей',
        '   {{ 30 | прописью }} → тридцать',
        '   {{ "27.05.2026" | дата_рус }} → 27 мая 2026 г.',
        '',
        'ПАДЕЖИ: им, род, дат, вин, тв, пр',
    ]
    for i, line in enumerate(info_lines, start=1):
        info.cell(row=i, column=1, value=line)
        if i == 1:
            info.cell(row=i, column=1).font = openpyxl.styles.Font(bold=True, size=14)
    info.column_dimensions['A'].width = 90

    Path(registry_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(registry_path)
